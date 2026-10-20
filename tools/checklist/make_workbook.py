"""Rebuild Harmony's compliance calendar in the Avaana checklist layout.

Brittany asked for one checklist in the shape of Avaana's 2023 NDIS application
checklist: module, Practice Standard, outcome, indicators, then the documents.
This takes the 16 Oct 2026 compliance calendar (every audited document, its link
and its review dates) and replaces its flat 'All documents' tab with a
'Checklist' tab in that shape, adding version columns and highlighting every
form someone must fill in by whose file it belongs in. The other tabs are kept.

It is a one-off: once Harmony types review dates into the result, that workbook
is the master, and export_web.py reads it for the website page.

    python3 make_workbook.py BASE.xlsx OUT.xlsx
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from send_lists import SEND_TABS

HERE = Path(__file__).parent
STANDARDS = json.loads((HERE / "standards.json").read_text())
# The lockup cropped, unchanged, from the letterhead band the branded set is printed on.
LOGO = HERE.parents[1] / "site" / "img" / "hc-lockup.png"  # one copy, shared with the website
LETTERHEAD_CREAM = "FDFBF6"

# Harmony Care Palette 01 (site/index.html tokens).
FOREST, SAGE, ROSE, CREAM, INK, MUTED, RULE = (
    "38553E", "9EAE8F", "F17990", "F3E3CA", "2E2A26", "5F564C", "DFD5C4")
TINT = {"Participant": "FBE0E6", "Worker": "DCE7D5", "Organisation": CREAM}  # rose and sage, lightened
FILE_LABEL = {"Participant": "Each participant", "Worker": "Each worker",
              "Organisation": "Organisation"}

MODULES = [  # (module_ord in the Guidelines, label)
    (1, "Core Module"),
    (2, "Module 1: High intensity daily personal activities"),
    (4, "Module 2A: Implementing behaviour support plans"),
    (6, "Module 4: Specialised support coordination"),
    (8, "Module 5A: Supported independent living"),
]
# Groups that are not a single Guidelines section.
EXTRA = {
    "M1-ALL": (2, "Every Module 1 support",
               "These documents serve every high intensity support Harmony delivers "
               "(bowel care, enteral feeding, urinary catheter and wound management)."),
    "M5A-ALL": (8, "Every SIL standard",
                "These documents serve every part of Module 5A."),
    "BUS": (9, "Business and registration evidence",
            "Not a Practice Standard. The papers the Commission and the auditor ask to see: "
            "business plan, insurances, registrations and the application answers."),
}

AREA_TO_SECTION = {
    "Rights: Person-centred supports": "6",
    "Rights: Privacy and dignity": "8",
    "Rights: Independence and informed choice": "9",
    "Rights: Violence, abuse, neglect, exploitation and discrimination": "10",
    "Governance: Governance and operational management": "11",
    "Governance: Risk management": "12",
    "Governance: Quality management": "13",
    "Governance: Information management": "14",
    "Governance: Feedback and complaints": "15",
    "Governance: Incident management": "16",
    "Governance: Human resource management": "17",
    "Governance: Continuity of supports": "18",
    "Governance: Emergency and disaster management": "18A",
    "Provision: Access to supports": "19",
    "Provision: Support planning": "20",
    "Provision: Service agreements": "21",
    "Provision: Responsive support provision": "22",
    "Provision: Transitions to or from the provider": "23",
    "Environment: Safe environment": "24",
    "Environment: Infection prevention and control": "24",  # IPC is indicators 7 to 11 of s 24
    "Environment: Client money and property": "25",
    "Environment: Management of medication": "26",
    "Environment: Mealtime management": "26A",
    "Environment: Management of waste": "27",
    "Business and registration": "BUS",
}
# The module areas hold several standards each, so place those documents by title.
# First match wins; every module document must match something (checked below).
TITLE_TO_SECTION = [
    ("Module 1", r"bowel", "29"),
    ("Module 1", r"enteral|tube", "30"),
    ("Module 1", r"catheter", "32"),
    ("Module 1", r"wound", "35"),
    ("Module 1", r".", "M1-ALL"),
    ("Module 2A", r"^Implementing Behaviour Supports Plans Policy|Internal Audit Schedule", "47"),
    ("Module 2A", r"Template Letter to PBSP", "49"),
    ("Module 2A", r"\(Interim\)", "54"),
    ("Module 2A", r"^Restrictive Practices Register", "51"),
    ("Module 2A", r"^Review Register|ABC Data", "52"),
    ("Module 2A", r".", "50"),
    ("Module 4", r"Conflict of Interest", "66"),
    ("Module 4", r"Services Agreement", "65"),
    ("Module 4", r".", "64"),
    ("Module 5a", r"Supported Decision-Making", "72B"),
    ("Module 5a", r"Safeguarding|Co-Tenancy Concern", "72C"),
    ("Module 5a", r"Practice Governance|Emergency Plan|Vision, Values|Training Plan|Induction to a House"
                  r"|Co-Tenant Consultation|House Meeting|Competency Assessment", "72D"),
    ("Module 5a", r"Tenancy|Transition or Exit|Services Agreement|Schedule to the Services|Collaboration"
                  r"|Conflict of Interest", "72E"),
    ("Module 5a", r".", "M5A-ALL"),
]
# Placed by the indicator that names them (checked 19 Oct 2026 against the Guidelines).
TITLE_OVERRIDE = {
    # filed under person-centred supports, but the policy for 'Individual values and beliefs'
    "Diversity Policy": "7",
    "SDA and SIL Collaboration Agreement": "21",  # 21(4): SIL delivered in SDA dwellings
    "Business Plan": "11",  # 11(4): strategic and business planning
    "Professional Indemnity and Public Liability certificate of currency": "12",  # 12(6): insurance
    "Vehicle insurance certificate (comprehensive, transport vehicle)": "12",
    "WorkCover Queensland certificate of currency": "12",
    "Internal Audit Schedule (Module 2A)": "13",  # 13(2): the internal audit program
    "Subject to a Significant Risk Factor Register": "20",  # 20(2A): reliance on the provider
    "Statement of support for the reduction and elimination of the use of restrictive practice": "47",  # 47(3)
    # both training records serve every Module 1 support ('wound' in the new titles must not file them under 35)
    "High Intensity Worker Training Plan and Competency Record (with wound care)": "M1-ALL",
    "High Intensity Worker Training Plan and Competency Record (without wound care)": "M1-ALL",
}
# Standards no document is filed under, but which these documents also meet.
SEE_ALSO = {
    "48": ["Restrictive Practices Register (plans with a regulated restrictive practice)"],  # 48(3)
    "53": ["Incident Report Form", "Incident Management Register"],  # 53(3)
    "17": ["Position Description: Support Coordinator"],
    "21": ["Support Coordination Services Agreement"],
}
# Two pairs of files share a title; their own approval blocks tell them apart.
RENAME_BY_ID = {
    "1W6DVKrS-vxcNGVZTdUk6GVqvmCJJYvoP": "Exit and Transition Form (draft copy, Module 1 and SIL set)",
    "16JGOsNDrYd_Zt0uiUFbtSOXjhqQU7T2y":
        "High Intensity Worker Training Plan and Competency Record (with wound care)",
    "1kVLrYz8JW_v3Q8CLwObqXBO6pMlc0ZXa":
        "High Intensity Worker Training Plan and Competency Record (without wound care)",
}
STATUS_BY_ID = {"1W6DVKrS-vxcNGVZTdUk6GVqvmCJJYvoP": "Draft"}  # its approval block reads 'Draft for'

# Where the 16 Oct reading was wrong, from the document itself (checked 19 Oct 2026).
PLAN_DEVELOPERS = ("The client (or guardian) and a worker develop it, after a qualified health "
                   "practitioner's assessment; the Principal checks it is complete")
CORRECTIONS = {
    "Risk Taking Form": ("completed_by", "Harmony (Key Management Personnel) documents it with the client",
                         "the form says Harmony documents the activity and its risks"),
    "Client Support Plan": ("completed_by", "Director (Principal) with the client and their support network; "
                                            "the client fills in the sections about themselves",
                            "the form asks the client to 'please fill the below'"),
    "Complex Bowel Care Support Plan": ("completed_by", PLAN_DEVELOPERS,
                                       "its 'Developed by' rows name the client and a worker"),
    "Enteral Feeding Support Plan": ("completed_by", PLAN_DEVELOPERS,
                                     "its 'Developed by' rows name the client and a worker"),
    "Urinary Catheter Support Plan": ("completed_by", PLAN_DEVELOPERS,
                                      "its 'Developed by' rows name the client and a worker"),
}
# Things only Harmony can fix in its own documents; added to 'Notes for the next audit'.
WOUND = ("Complex wound management is in the audited folder but was not part of the Module 1 scope "
         "audited on 6 October 2026 (bowel care, enteral feeding, urinary catheter).")
HARMONY_NOTES = {
    "Risk Taking Form": "Its Harmony signature line still names the sole trader (Victoria J Scilipote, "
                        "ABN 50 660 758 851); on the company's paper the party is Harmony Care Support "
                        "Services Pty Ltd, ABN 99 691 543 834.",
    "Client Support Plan": "Its Harmony signature line reads 'on behalf of Victoria J Scilipote' under the "
                           "company's ABN; the signing party should be the company.",
    "SDA and SIL Collaboration Agreement": "The execution line is pre-printed with 2022. The SDA provider "
                                           "also completes the Home Risk Assessment in the annexure.",
    "Position Description: Disability Support Worker": "Two versions of this position description are on "
                                                        "file (see the 'second version'); confirm which is current.",
    "Position Description: Disability Support Worker (second version)": "Two versions of this position "
                                                                        "description are on file; this one "
                                                                        "carries the company's ABN.",
    "Business Continuity Plan (PDF, July 2026)": "Two continuity plans are on file; this one is signed, the "
                                                 "Word management plan is not. Confirm which is current.",
    "Business Continuity, Emergency and Disaster Preparedness Management Plan":
        "Two continuity plans are on file; the July 2026 PDF is signed, this one is not. Confirm which is current.",
    "Exit Interview Form": "Its approval block still holds a template placeholder ('{{Entity_Name}}', July "
                           "2020) where the approver should be.",
    "Exit and Transition Form (draft copy, Module 1 and SIL set)": "Its approval block reads 'Draft for'. The "
                                                                   "core folder copy is the approved one.",
    "Complex Wound Management Policy": WOUND, "Complex Wound Management Support Plan": WOUND,
    "Wound Observation and Dressing Chart": WOUND, "Complex Wound Management (Easy Read)": WOUND,
}
VERSIONS = json.loads((HERE / "versions.json").read_text())  # from read_versions.py
# Harmony's Document Register as audited (Brittany's Documents folder): every row is version 1,
# approved by Victoria J Scilipote in October 2023, next review October 2025. It fills only what a
# document's own approval block and the 16 Oct calendar left empty. (The Rise version's copy of the
# register reads October 2026; the audited copy is the record, so it is the one used.)
DOCUMENT_REGISTER = {  # register entry -> the checklist titles it covers
    'Administering Medication Checklist & Medication Chart': ['Administering Medication Checklist and Medication Chart'],
    'Advocacy of Support Person Request Form': ['Advocacy or Support Person Request Form'],
    'Business Plan': ['Business Plan'],
    'Business Continuity and Disaster Management Plan': ['Business Continuity, Emergency and Disaster Preparedness Management Plan'],
    'Change of Supports': ['Change to Supports Agreement'],
    'Client Cash Reconciliation': ['Client Cash Reconciliation'],
    'Client Emergency Plan': ['Client Emergency Plan'],
    'Client Information Booklet': ['Client Information Booklet'],
    'Client Information Booklet (Easy Read)': ['Client Information Booklet (Easy Read)'],
    'Client Induction Checklist': ['Client Induction Checklist'],
    'Client Intake Form': ['Client Intake Form'],
    'Client Support Plan': ['Client Support Plan'],
    'Client Survey': ['Client Survey'],
    'Collection & Storage of Medicine Register': ['Collection and Storage of Medicine Register'],
    'Complaints Register': ['Complaints Register'],
    'Conflict of Interest Declaration': ['Conflict of Interest Declaration (support coordination)'],
    'Conflict of Interest Register': ['Conflict of Interest Register'],
    'Continuous Improvement Register': ['Continuous Improvement Register'],
    'Delegation of Authority': ['Delegation of Authority'],
    'Emergency Waste Management Plan': ['Emergency Waste Management Plan'],
    'Emergency Test Register': ['Emergency Test Register'],
    'Exit and Transition Form': ['Exit and Transition Form'],
    'Exit Interview Form': ['Exit Interview Form'],
    'Feedback and Complaints Assessment, Investigation and Resolution Considerations': ['Feedback and Complaints Assessment, Investigation and Resolution Considerations'],
    'Feedback and Complaints Criteria for Complaint Manager or Incident Manager': ['Criteria for Complaint Manager or Incident Manager'],
    'Feedback and Complaints Form': ['Feedback and Complaints Form'],
    'Feedback and Complaints Procedural Fairness Considerations': ['Feedback and Complaints Procedural Fairness Considerations'],
    'Feedback and Complaints Summary': ['Feedback and Complaints Summary'],
    'Feedback and Compliments Register': ['Feedback and Compliments Register'],
    'HR Performance Appraisal': ['HR Performance Appraisal'],
    'Incident Management Register': ['Incident Management Register'],
    'Incident Report Form': ['Incident Report Form'],
    'Internal Audit Schedule': ['Internal Audit Schedule (Core)'],
    'Interview Report Pro Forma': ['Interview Report Pro Forma (Recruitment and Selection Report)'],
    'Legislation Register': ['Legislation Register'],
    'Management Meeting Agenda': ['Management Meeting Agenda and Minutes'],
    'Managing and Reducing Known Risks Matrix': ['Managing and Reducing Known Risks Matrix'],
    'Mealtime Management Care Plan': ['Mealtime Management Plan'],
    'Money and Property Consent Form': ['Money and Property Consent Form'],
    'My NDIS Support Record': ['Worker Timesheets (NDIS Support Record)'],
    'Policy Register': ['Policy Register'],
    'Position Descriptions': ['Position Description: Principal', 'Position Description: Disability Support Worker', 'Position Description: Disability Support Worker (Employment Consultant)', 'Position Description: Disability Support Worker (second version)', 'Position Description: Disability Support Worker (Travel and Transport)', 'Position Description: Support Coordinator'],
    'Potential Staff Reference Check Form': ['Potential Staff Reference Check Form'],
    'Pre Employment Collection Form': ['Pre Employment Collection Form'],
    'Privacy Consent Form': ['Privacy Consent Form'],
    'Progress Notes Template': ['Progress Notes Template'],
    'Risk Assessment Form': ['Risk Assessment Form'],
    'Risk Assessed Role Register': ['Risk Assessed Role Register (Human Resources Risk Assessed Role Register)'],
    'Risk Taking Form': ['Risk Taking Form'],
    'Risk Register': ['Risk Register'],
    'SDA and SIL Collaboration Agreement': ['SDA and SIL Collaboration Agreement'],
    'Service Schedule': ['Services and Supports Schedule'],
    'Services Agreement': ['Services Agreement'],
    'SIL Services Agreement': ['Services Agreement (SIL)'],
    'Support Coordination Services Agreement': ['Support Coordination Services Agreement'],
    'Services Medication and Participant Details': ['Services Medication and Participant Details'],
    'Staff Exit Checklist': ['Staff Exit Checklist'],
    'Staff File Checklist': ['Staff File Checklist'],
    'Staff Induction Checklist': ['Staff Induction Checklist'],
    'Staff Performance Improvement Plan': ['Staff Performance Improvement Plan'],
    'Subject to a Significant Risk Factor': ['Subject to a Significant Risk Factor Register'],
    'Training and Development Register': ['Training and Development Register'],
    'Training Feedback Form': ['Training Feedback Form'],
    'Waste Management Plan': ['Waste Management Plan'],
    'Waste Management Report': ['Waste Management Report'],
    'Waste Removal Records Register': ['Waste Removal Records Register'],
    'WHS Walk-Around Risk Assessment Checklist': ['Work Health and Home Safety Risk Assessment Checklist'],
    'Worker Declarations': ['Worker Declarations'],
    'Worker Supervision Record': ['Worker Supervision Record'],
    'ABC Data Form': ['ABC Data Form'],
    'Behaviour Support Plan': ['Behaviour Support Plan (Comprehensive)'],
    'Coaching Form for PBSP': ['Coaching Form for Positive Behaviour Support Plan'],
    'Disability Support Worker Position Description (IBS)': ['Position Description: Disability Support Worker (Implementing Behaviour Supports)'],
    'Interim Behaviour Support Plan': ['Behaviour Support Plan (Interim)'],
    'Register for PBSP that contain a regulated restrictive practice': ['Restrictive Practices Register (plans with a regulated restrictive practice)'],
    'Review Register for PBSP that contain a regulated restrictive practice': ['Review Register for Behaviour Support Plans with a Regulated Restrictive Practice'],
    'Statement of support for the reduction and elimination of the use of restrictive practice': ['Statement of support for the reduction and elimination of the use of restrictive practice'],
    'Template Letter to SBSP': ['Template Letter to PBSP (behaviour support practitioner)'],
}
REGISTER_ROW = {"version": "1", "approved_by": "Victoria J Scilipote", "approval_date": "October 2023"}
REGISTER_REVIEW = datetime.datetime(2025, 10, 31)

# Documents Harmony should hold that no audited folder has. Shown with the status 'Missing'.
MISSING_DOCS = [
    {"title": "Statement of support for the reduction and elimination of the use of restrictive practice",
     "area": "Module 2A: Implementing behaviour support plans", "type": "Statement", "for": "Organisation",
     "when": "In force at all times: Harmony's stated commitment to reducing and eliminating restrictive "
             "practices (indicator 47(3)).",
     "how_often": "Review by the date in the Document Register (October 2025).",
     "completed_by": "Principal", "signed_by": "Principal", "kept_where": "No file found",
     "next_review": REGISTER_REVIEW, "review_months": 24,
     "rule": "Document Register: version 1, approved October 2023, next review October 2025",
     "notes": "Listed in Harmony's Document Register, but no file was found in the audited folders or the Rise "
              "version. Find it, or write it before the next audit."},
    {"title": "NDIS certificate of registration", "area": "Business and registration",
     "type": "Business paper", "for": "Organisation",
     "when": "Held at all times. Shows Harmony's registration groups and the date registration ends.",
     "how_often": "Renew before the end date on the certificate (see 'Rules and sources', item 13).",
     "completed_by": "Issued by the NDIS Quality and Safeguards Commission", "signed_by": "Not signed",
     "kept_where": "No file found", "next_review": None, "review_months": None,
     "rule": "'Rules and sources', item 13",
     "notes": "Not in the audited folders: the 'RegistrationCertificate.pdf' there is a Queensland vehicle "
              "registration. File the Commission's certificate here so the registration end date can be tracked."},
]
for _d in MISSING_DOCS:
    _d.update(id=None, url=None, last_reviewed=None, fixed_status="Missing")


FILLED_IN = {"Form", "Register", "Plan", "Record", "Checklist", "Chart", "Agreement", "Template"}
FILE_OF = {"Each client": "Participant", "Each behaviour support plan": "Participant",
           "Each worker": "Worker", "Each new worker": "Worker",
           "Organisation": "Organisation", "Each SIL house": "Organisation",
           "Each incident": "Organisation", "Each complaint": "Organisation"}
# A template the organisation fills in to make a document, not a worker's own form.
FILE_OVERRIDE = {"Position Description Template": "Organisation"}

# Harmony's Policy Register (edited 28 Aug 2026): register title -> our title.
POLICY_REGISTER = {
    "Choice Advocacy and Control Policy": "Choice, Advocacy and Control Policy",
    "Diversity Policy": "Diversity Policy",
    "Preventing and Responding to Violence, Abuse, Neglect, Exploitation and Discrimination Policy":
        "Preventing and Responding to Violence, Abuse, Neglect, Exploitation and Discrimination Policy",
    "Privacy Policy": "Privacy and Dignity Policy",
    "Promoting and Protecting Rights Policy": "Promoting and Protecting Rights Policy",
    "Feedback and Complaints Management Policy": "Feedback and Complaints Management Policy",
    "Incident Management and Reporting Policy": "Incident Management and Reporting Policy",
    "Quality Management and Continuous Improvement Policy": "Quality Management and Continuous Improvement Policy",
    "Risk Management Policy": "Risk Management Policy",
    "Emergency and Disaster Preparedness Policy": "Emergency and Disaster Preparedness Policy",
    "Work Health and Safety Policy": "Work Health and Safety Policy",
    "COVID-19 Pandemic Management Policy": "COVID-19 Pandemic Management Policy",
    "Governance Policy": "Governance Policy",
    "Human Resources Management Policy": "Human Resources Management Policy",
    "Information Management Policy": "Information Management Policy",
    "Conflict of Interest Policy": "Conflict of Interest Policy",
    "Provision of Supports Policy": "Provision of Supports Policy",
    "Client Money and Property Policy": "Client Money and Property Policy",
    "Infection Control Policy": "Infection Control Policy",
    "Support Coordination Policy": "Support Coordination Policy",
    "Supported Independent Living Policy": "Supported Independent Living Policy (Version 1, core folder copy)",
    "Mealtime Management Policy": "Mealtime Management Policy",
    "Management of Medication Policy": "Management of Medication Policy",
    "Waste Management Policy": "Waste Management Policy",
    "Implementing Behaviour Supports Plans Policy & Procedures":
        "Implementing Behaviour Supports Plans Policy and Procedures",
}
VERSION_KNOWN = {"Supported Independent Living Policy": "2"}  # replaces Version 1 (6 Oct 2026 pack)

COLS = [  # (header, width, key)
    ("No.", 6, "no"),
    ("Module", 16, "module"),
    ("Practice Standard", 24, "standard"),
    ("Document (opens the file)", 42, "title"),
    ("Type", 13, "type"),
    ("Filled in for", 14, "file"),
    ("When it is needed", 40, "when"),
    ("Completed by", 22, "completed_by"),
    ("Signed by", 26, "signed_by"),
    ("Version", 9, "version"),
    ("Approved by", 17, "approved_by"),
    ("Approval date", 13, "approval_date"),
    ("Next review on record", 13, "next_review"),
    ("Review every (months)", 10, "review_months"),
    ("Last reviewed (type the date)", 14, "last_reviewed"),
    ("Next due", 12, None),
    ("Status", 18, None),
    ("Notes for the next audit", 80, "notes"),
    ("How often", 40, "how_often"),
    ("Kept where", 24, "kept_where"),
    ("Harmony's rule (source)", 46, "rule"),
    ("Section", 8, "section"),  # hidden: the Guidelines section, read by export_web.py
]
COL = {key or head: i + 1 for i, (head, _, key) in enumerate(COLS)}
TYPED = {"version", "approved_by", "approval_date", "last_reviewed"}
HEADER_ROW = 5  # row 1 carries the logo

thin = Side(style="thin", color=RULE)
BORDER = Border(bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
F = lambda size=10, **k: Font(name="Arial", size=size, **k)  # noqa: E731


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def add_logo(ws, last_col, height_px=62):
    """The Harmony Care lockup across row 1, on the letterhead's own cream."""
    img = XLImage(str(LOGO))
    img.width, img.height = round(img.width * height_px / img.height), height_px
    ws.add_image(img, "A1")
    ws.row_dimensions[1].height = height_px * 0.75 + 10
    for i in range(1, last_col + 1):
        ws.cell(1, i).fill = fill(LETTERHEAD_CREAM)


def read_base(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["All documents"]
    head = [c.value for c in ws[1]]
    docs = []
    for row in ws.iter_rows(min_row=2):
        if not row[1].value:
            continue
        d = dict(zip(head, (c.value for c in row)))
        status = d["Status"]
        url = row[1].hyperlink.target if row[1].hyperlink else None
        fid = url.split("/d/")[1].split("/")[0] if url else None
        docs.append({
            "id": fid,
            "area": d["Area"],
            "title": RENAME_BY_ID.get(fid, d["Document (opens the file)"]),
            "url": url,
            "type": d["Type"], "for": d["For"],
            "when": d["When it is needed"], "how_often": d["How often"],
            "completed_by": d["Completed by"], "signed_by": d["Signed by"],
            "kept_where": d["Kept where"],
            "next_review": d["Next review on record"],
            "review_months": d["Review every (months)"],
            "last_reviewed": d["Last reviewed (type the date)"],
            "rule": d["Harmony's rule (source)"], "notes": d["Notes for the next audit"],
            # A few rows carry a fixed status (working copy, superseded) instead of the formula.
            "fixed_status": STATUS_BY_ID.get(fid) or (
                None if (status is None or str(status).startswith("=")) else status),
        })
    return wb, docs


def place(doc):
    if doc["title"] in TITLE_OVERRIDE:
        return TITLE_OVERRIDE[doc["title"]]
    if doc["area"] in AREA_TO_SECTION:
        return AREA_TO_SECTION[doc["area"]]
    area = doc["area"].split(":")[0]
    for a, pattern, section in TITLE_TO_SECTION:
        if a == area and re.search(pattern, doc["title"], re.I):
            return section
    raise SystemExit(f"no Practice Standard for {doc['area']!r} / {doc['title']!r}")


def _add(notes, line):
    return f"{notes} {line}" if notes else line


def enrich(docs):
    by_title = {d["title"]: d for d in docs}
    for reg_title, ours in POLICY_REGISTER.items():
        d = by_title.get(ours)
        if d is None:
            raise SystemExit(f"Policy Register entry {reg_title!r} matches no document ({ours!r})")
        d.update(version="1", approved_by="Victoria J Scilipote", approval_date="October 2023")
    for title, v in VERSION_KNOWN.items():
        by_title[title]["version"] = v
    for d in docs:
        own = VERSIONS.get(d["id"], {})  # the document's own approval block wins
        if own.get("version"):
            d["version"] = own["version"]
        who = own.get("approved_by", "")
        if who.startswith("Victoria J Scilipote"):
            d["approved_by"] = "Victoria J Scilipote"
        if own.get("approval_date"):
            d["approval_date"] = own["approval_date"]
    for entry, titles in DOCUMENT_REGISTER.items():
        for title in titles:
            matches = [d for d in docs if d["title"] == title]
            if not matches:
                raise SystemExit(f"Document Register entry {entry!r} names no document {title!r}")
            for d in matches:
                for key, value in REGISTER_ROW.items():
                    if not d.get(key):
                        d[key] = value
                if not d["next_review"]:
                    d["next_review"] = REGISTER_REVIEW
                    d["notes"] = _add(d["notes"], f"Review date from Harmony's Document Register "
                                                  f"('{entry}': October 2025).")
    for title, (field, value, why) in CORRECTIONS.items():
        d = by_title[title]
        d[field] = value
        d["notes"] = _add(d["notes"], f"Corrected 19 Oct 2026 from the document: {why}.")
    for title, note in HARMONY_NOTES.items():
        for d in docs:
            if d["title"] == title:
                d["notes"] = _add(d["notes"], note)
    missing = [t for t in list(CORRECTIONS) + list(HARMONY_NOTES) if t not in by_title]
    if missing:
        raise SystemExit(f"corrections name documents that do not exist: {missing}")
    for d in docs:
        d["section"] = place(d)
        if d["type"] in FILLED_IN:
            d["file"] = FILE_OVERRIDE.get(d["title"]) or FILE_OF[d["for"]]
        else:
            d["file"] = None


def groups(docs):
    """(module label, section key, heading, outcome, rows, clause, docs, also) in Guidelines order.

    'also' are documents filed under another standard that meet this one too."""
    sections = {s["section_no"]: s for s in STANDARDS["sections"]}
    order = []
    for ord_, label in MODULES + [(9, "Business and registration")]:
        keys = [k for k, (m, _, _) in EXTRA.items() if m == ord_ and k != "BUS"]
        keys += [s["section_no"] for s in STANDARDS["sections"] if s["module_ord"] == ord_]
        keys += ["BUS"] if ord_ == 9 else []
        for k in keys:
            mine = [d for d in docs if d["section"] == k]
            also = [d for t in SEE_ALSO.get(k, []) for d in docs if d["title"] == t]
            if len(also) != len(SEE_ALSO.get(k, [])):
                raise SystemExit(f"SEE_ALSO for {k} names a document that does not exist")
            if not mine and not also:
                continue
            if k in EXTRA:
                _, heading, outcome = EXTRA[k]
                order.append((label, k, heading, outcome, [], None, mine, also))
            else:
                s = sections[k]
                order.append((label, k, s["heading"], s["outcome"], s["rows"],
                              f"Practice Standards Rules, {s['rules_ref']}; "
                              f"Quality Indicators Guidelines, section {k}", mine, also))
    placed = sum(len(g[6]) for g in order)
    if placed != len(docs):
        raise SystemExit(f"{len(docs) - placed} documents fell outside every group")
    return order


def indicator_lines(rows):
    """One string per top-level indicator, its sub-items on their own lines."""
    out = []
    for i, r in enumerate(rows):
        if r["parent"] is None:
            out.append([f"{r['item']}. {r['text']}"])
        else:
            out[-1].append(f"     ({r['item']}) {r['text']}")
        for c in r.get("continued") or []:
            out[-1].append(f"     {c}")
    return ["\n".join(x) for x in out]


def height(text, width_units, size=10):
    """Row height in points that fits wrapped text in a cell this many units wide."""
    per_line = max(1, int(width_units * 1.0 * 10 / size))
    lines = sum(max(1, -(-len(p) // per_line)) for p in str(text).split("\n"))
    return min(409, max(15, lines * size * 1.32 + 4))


def write_checklist(wb, order, at):
    ws = wb.create_sheet("Checklist", at)
    for i, (_, w, _) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions[get_column_letter(COL["section"])].hidden = True
    last = get_column_letter(len(COLS))
    span = sum(w for _, w, k in COLS[COL["title"] - 1:COL["signed_by"]])
    d0, d1 = get_column_letter(COL["title"]), get_column_letter(COL["signed_by"])

    add_logo(ws, len(COLS) - 1)
    ws["A2"] = "Documents and forms by Practice Standard"
    ws["A2"].font = F(bold=True, size=14, color=FOREST)
    ws["A3"] = "Highlighted rows are forms someone must fill in, coloured by whose file they belong in:"
    ws["A3"].font = F(color=MUTED)
    for col, who in zip(("F", "G", "H"), ("Participant", "Worker", "Organisation")):
        ws[f"{col}3"] = FILE_LABEL[who] if who != "Organisation" else "The organisation"
        ws[f"{col}3"].fill = fill(TINT[who])
        ws[f"{col}3"].font = F(bold=True)
        ws[f"{col}3"].alignment = Alignment(horizontal="center")
    ws["I3"] = "No colour: policies and guides, read but not filled in"
    ws["I3"].font = F(color=MUTED)
    ws["A4"] = ("Wording of each outcome and indicator: NDIS Practice Standards Rules (compilation 6) and "
                "Quality Indicators Guidelines (compilation 3), both in force 1 July 2026. The 1 and 2 "
                "buttons at the top left fold the indicators away and back. Wording quoted unchanged from the "
                "Federal Register of Legislation (legislation.gov.au, F2026C00527 and F2026C00528).")
    ws["A4"].font = F(color=MUTED, italic=True)

    for i, (head, _, key) in enumerate(COLS, 1):
        c = ws.cell(HEADER_ROW, i, head)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = fill(SAGE if key in TYPED else FOREST)
        if key in TYPED:
            c.font = F(bold=True, color=INK)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[HEADER_ROW].height = 44

    r = HEADER_ROW + 1
    no = 0
    module_seen = None
    doc_rows = []
    # Number every document first, so a standard can point at one filed further down.
    for module, _, heading, *_, docs, _also in order:
        for d in docs:
            no += 1
            d["no"], d["module"], d["standard"] = no, module.split(":")[0], heading
    for module, key, heading, outcome, rows, clause, docs, also in order:
        if module != module_seen:
            module_seen = module
            ws.merge_cells(f"A{r}:{d1}{r}")
            c = ws.cell(r, 1, module.upper())
            c.font = F(bold=True, size=12, color="FFFFFF")
            c.fill = fill(FOREST)
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[r].height = 22
            r += 1
        # The standard: its name, then its outcome across the document columns.
        ws.cell(r, COL["standard"], heading).font = F(bold=True, color=FOREST, size=11)
        ws.cell(r, COL["standard"]).alignment = WRAP
        ws.merge_cells(f"{d0}{r}:{d1}{r}")
        text = f"Outcome: {outcome}" + (f"\n({clause})" if clause else "")
        c = ws.cell(r, COL["title"], text)
        c.font = F(italic=True, color=INK)
        c.alignment = WRAP
        for i in range(1, len(COLS)):
            ws.cell(r, i).border = Border(top=Side(style="medium", color=FOREST))
        ws.row_dimensions[r].height = max(height(heading, 24, 11), height(text, span))
        r += 1
        lines = indicator_lines(rows)
        first = r
        for n, line in enumerate(lines):
            ws.cell(r, COL["standard"], "Indicators to be demonstrated" if n == 0 else None).font = F(color=MUTED, size=9)
            ws.merge_cells(f"{d0}{r}:{d1}{r}")
            c = ws.cell(r, COL["title"], line)
            c.font = F(color=MUTED, size=9)
            c.alignment = WRAP
            ws.row_dimensions[r].height = height(line, span, 9)
            r += 1
        if lines:
            ws.row_dimensions.group(first, r - 1, outline_level=1, hidden=False)
        for d in also:
            ws.cell(r, COL["standard"], "Also meets this standard").font = F(color=MUTED, size=9)
            ws.merge_cells(f"{d0}{r}:{d1}{r}")
            c = ws.cell(r, COL["title"], f"No. {d['no']}, {d['title']} (filed under {d['standard']})")
            c.font = F(color=FOREST, underline="single", italic=True)
            c.alignment = WRAP
            if d["url"]:
                c.hyperlink = d["url"]
            r += 1
        for d in docs:
            for i, (_, w, k) in enumerate(COLS, 1):
                if k is None:
                    continue
                v = d.get(k)
                c = ws.cell(r, i, v)
                c.font = F()
                c.alignment = WRAP
                c.border = BORDER
                if k == "file":
                    c.value = FILE_LABEL[v] if v else "Not filled in"
                    c.font = F(bold=bool(v), color=INK if v else MUTED)
                if k in ("next_review", "last_reviewed") and v:
                    c.number_format = "d mmm yyyy"
            t = ws.cell(r, COL["title"])
            if d["url"]:
                t.hyperlink = d["url"]
                t.font = F(color=FOREST, underline="single", bold=bool(d["file"]))
            L = lambda k: f"{get_column_letter(COL[k])}{r}"  # noqa: E731
            due = ws.cell(r, COL["Next due"],
                          f'=IF({L("last_reviewed")}<>"",IF({L("review_months")}="","",'
                          f'EDATE({L("last_reviewed")},{L("review_months")})),'
                          f'IF({L("next_review")}="","",{L("next_review")}))')
            due.number_format = "d mmm yyyy"
            due.font, due.border, due.alignment = F(), BORDER, WRAP
            st = ws.cell(r, COL["Status"], d["fixed_status"] or (
                f'=IF({L("Next due")}="","No review date on record",IF({L("Next due")}<TODAY(),"Overdue",'
                f'IF({L("Next due")}<TODAY()+30,"Due within 30 days","In date")))'))
            st.font, st.border, st.alignment = F(bold=True), BORDER, WRAP
            if d["file"]:
                for i in range(1, len(COLS)):
                    ws.cell(r, i).fill = fill(TINT[d["file"]])
            # No fixed height: Excel, Sheets and LibreOffice fit an unmerged row to its wrapped text.
            doc_rows.append(r)
            r += 1

    end = r - 1
    st = get_column_letter(COL["Status"])
    due = get_column_letter(COL["Next due"])
    red, amber, green, grey = fill("F4C7C3"), fill("FCE8B2"), fill("D7EBD5"), fill("E8E4E0")
    rng = f"{st}{HEADER_ROW + 1}:{st}{end}"
    for value, f_ in (("Overdue", red), ("Due within 30 days", amber), ("No review date on record", amber),
                      ("In date", green), ("Superseded", grey), ("Working copy", grey), ("Draft", grey),
                      ("Missing", red)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{st}{HEADER_ROW + 1}="{value}"'], fill=f_))
    rng = f"{due}{HEADER_ROW + 1}:{due}{end}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({due}{HEADER_ROW + 1}),{due}{HEADER_ROW + 1}<TODAY())"], fill=red))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({due}{HEADER_ROW + 1}),{due}{HEADER_ROW + 1}<TODAY()+30)"], fill=amber))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last}{end}"
    ws.freeze_panes = f"{get_column_letter(COL['type'])}{HEADER_ROW + 1}"
    ws.sheet_properties.outlinePr.summaryBelow = False
    # Printed: A3 landscape, one page wide, the checklist columns only (No. to Status).
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.print_area = f"A1:{st}{end}"
    return ws, end


SEND_COLS = [("No.", 6), ("Form (opens the file)", 46), ("They need to", 30), ("When it is needed", 44),
             ("Who fills it in", 32), ("Who signs", 32), ("Practice Standard", 30)]
NOT_SENT = ("Superseded", "Working copy", "Draft", "Missing")


def _send_sheet(wb, at, name, title, intro, tint, tab_colour, groups):
    """One tab: a title, then each group's forms under a band, every title linked."""
    ws = wb.create_sheet(name, at)
    ws.sheet_properties.tabColor = tab_colour
    for i, (_, w) in enumerate(SEND_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = get_column_letter(len(SEND_COLS))
    add_logo(ws, len(SEND_COLS))
    ws["A2"] = title
    ws["A2"].font = F(bold=True, size=14, color=FOREST)
    ws["A3"] = intro
    ws["A3"].font = F(color=MUTED)
    for i, (head, _) in enumerate(SEND_COLS, 1):
        c = ws.cell(HEADER_ROW, i, head)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = fill(FOREST)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[HEADER_ROW].height = 22
    r = HEADER_ROW + 1
    for group, rows in groups:
        ws.merge_cells(f"A{r}:{last}{r}")
        c = ws.cell(r, 1, group)
        c.font = F(bold=True, size=11, color=FOREST)
        c.border = Border(top=Side(style="medium", color=FOREST))
        ws.row_dimensions[r].height = 24
        c.alignment = Alignment(vertical="bottom")
        r += 1
        for d, action in rows:
            values = [d["no"], d["title"], action, d["when"], d["completed_by"], d["signed_by"],
                      f"{d['module']}: {d['standard']}"]
            for i, v in enumerate(values, 1):
                c = ws.cell(r, i, v)
                c.font = F(bold=(i == 3))
                c.alignment = WRAP
                c.border = BORDER
                if tint:
                    c.fill = fill(tint)
            t = ws.cell(r, 2)
            if d["url"]:
                t.hyperlink = d["url"]
                t.font = F(color=FOREST, underline="single", bold=True)
            r += 1
    ws.freeze_panes = f"C{HEADER_ROW + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"


KEPT_GROUPS = (("Participant", "Kept in each participant's file"), ("Worker", "Kept in each worker's file"),
               ("Organisation", "The organisation's own records"))


def destinations(docs):
    """Where every form goes. The Excel tabs, the PDF and the website all read this one answer.

    Returns ([(send tab, [(group, [(doc, action)])])], [(group, [(doc, action)])]): each 'Send to'
    list resolved to documents, then the forms on none of them, which Harmony fills in itself."""
    by_title = {}
    for d in docs:
        if d["fixed_status"] not in NOT_SENT:
            by_title.setdefault(d["title"], []).append(d)
    sent, out = set(), []
    for tab in SEND_TABS:
        groups = []
        for group, entries in tab["groups"]:
            rows = []
            for title, action in entries:
                if title not in by_title:
                    raise SystemExit(f"{tab['tab']}: no document called {title!r}")
                rows += [(d, action) for d in by_title[title]]
                sent.add(title)
            groups.append((group, rows))
        out.append((tab, groups))
    kept = [(label, [(d, "Harmony fills it in") for d in docs
                     if d["file"] == f and d["title"] not in sent and d["fixed_status"] not in NOT_SENT])
            for f, label in KEPT_GROUPS]
    return out, kept


def write_send_tabs(wb, docs, at):
    """The 'Send to' tabs, then 'Harmony fills in' for every form on none of them."""
    sent, kept = destinations(docs)
    for i, (tab, groups) in enumerate(sent):
        tint = TINT.get(tab["colour"]) if tab["colour"] != "Organisation" else None
        _send_sheet(wb, at + i, tab["tab"], tab["title"], tab["intro"], tint,
                        {"Participant": ROSE, "Worker": SAGE}.get(tab["colour"], FOREST), groups)
        print(f"  {tab['tab']}: {sum(len(g[1]) for g in groups)} rows")
    groups = kept
    _send_sheet(wb, at + len(SEND_TABS), "Harmony fills in", "Harmony fills these in itself (not sent out)",
                "Registers, records and plans the Principal, Director or Key Management Personnel complete and "
                "keep. Nobody outside Harmony needs to be sent these.", TINT["Organisation"], CREAM, groups)
    print(f"  Harmony fills in: {sum(len(g[1]) for g in groups)} rows")


# One row per form sent to one person. The website's "Copy rows for the tracker" button gives
# rows in exactly these first six columns, so they paste straight in. The same fields are what
# Rise will hold for each send, so the rows can be imported when Harmony moves to it.
TRACKER_COLS = [("Date sent", 13), ("Sent to", 24), ("Who they are", 20), ("Email", 30), ("Form", 44),
                ("They need to", 34), ("Returned on", 13), ("Approved by", 20), ("Approved on", 13),
                ("Status", 30), ("Notes", 40)]
TRACKER_ROWS = 400
ROLES = ["Participant", "Worker", "Practitioner or partner"]


def write_tracker(wb, at):
    ws = wb.create_sheet("Sent and returned", at)
    ws.sheet_properties.tabColor = FOREST
    for i, (_, w) in enumerate(TRACKER_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    add_logo(ws, len(TRACKER_COLS))
    ws["A2"] = "Sent and returned"
    ws["A2"].font = F(bold=True, size=14, color=FOREST)
    ws["A3"] = ("One row for each form sent to each person. Paste the rows the website's 'Copy rows for the "
                "tracker' button gives you, or type them. When the form comes back, type the date in "
                "'Returned on'; when it is checked, type who approved it and when. Status ticks itself.")
    ws["A3"].font = F(color=MUTED)
    head = 5
    for i, (name, _) in enumerate(TRACKER_COLS, 1):
        c = ws.cell(head, i, name)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = fill(FOREST)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[head].height = 24
    first, last = head + 1, head + TRACKER_ROWS
    typed = fill("FFF6DD")  # the staff file tab's colour for cells someone types in
    for r in range(first, last + 1):
        for col in range(1, len(TRACKER_COLS) + 1):
            c = ws.cell(r, col)
            c.font, c.border, c.alignment = F(), BORDER, Alignment(wrap_text=True, vertical="top")
            if col != 10:
                c.fill = typed
        for col in (1, 7, 9):
            ws.cell(r, col).number_format = "d mmm yyyy"
        ws.cell(r, 10).value = (f'=IF(E{r}="","",IF(I{r}<>"","\u2713 Done and approved",IF(G{r}<>"",'
                                f'"Returned: check and approve",IF(AND(ISNUMBER(A{r}),A{r}<TODAY()-14),'
                                f'"Chase up: sent over 2 weeks ago","Waiting for it to come back"))))')
        ws.cell(r, 10).font = F(bold=True)
    dv = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{first}:C{last}")
    rng = f"J{first}:J{last}"
    for text, colour in (("Done", "D7EBD5"), ("Returned", "FCE8B2"), ("Chase", "F4C7C3")):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",J{first}))'],
                                                       fill=fill(colour)))
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(TRACKER_COLS))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{head}:{head}"


def fix_start_here(wb, end):
    ws = wb["Start here"]
    ws["A1"] = None  # the lockup says the name
    add_logo(ws, 2)
    rv_row = None
    ws["A2"] = "Documents, forms and compliance calendar"
    st = get_column_letter(COL["Status"])
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "'All documents'!N2:N164" in c.value:
                c.value = c.value.replace("'All documents'!N2:N164", f"Checklist!{st}{HEADER_ROW + 1}:{st}{end}")
            elif c.value == "All documents":
                c.value = "Checklist"
                ws.cell(c.row, c.column + 1).value = (
                    "Laid out like the Avaana checklist: module, Practice Standard, its outcome and indicators, "
                    "then every document for it. Highlighted rows are forms someone must fill in: pink for each "
                    "participant's file, green for each worker's file, cream for the organisation's records. "
                    "Click a title to open the file. Filter by 'Filled in for' or 'Status'. The 'Send to' tabs "
                    "list what goes out to participants, workers and practitioners, and what each has to do; "
                    "'Harmony fills in' lists the forms that never leave the office. 'Sent and returned' "
                    "tracks each form sent out until it is back and approved.")
            elif isinstance(c.value, str) and c.value.startswith("After a review"):
                ws.cell(c.row, c.column + 1).value = (
                    "Type into the green-grey columns: the date in 'Last reviewed', and the version, approver and "
                    "approval date. Next due and Status update by themselves.")
            elif c.value == "Module 1 and SIL":
                rv_row = c.row + 1
            elif isinstance(c.value, str) and c.value.startswith("Built 16 October 2026"):
                c.value = c.value.replace("Built 16 October 2026",
                                          "Built 16 October 2026, set out by Practice Standard 19 October 2026,")
    # Two lines the 16 Oct sheet had no room for, written into its blank spacer rows.
    from copy import copy
    if rv_row is None or ws.cell(rv_row, 1).value is not None or ws.cell(27, 1).value is not None:
        raise SystemExit("Start here has changed; the spacer rows are no longer blank")
    for row, label, text in (
            (27, "Missing from the folders",
             "Two documents Harmony should hold are not in any folder: the 'Statement of support for the "
             "reduction and elimination of the use of restrictive practice' (listed in the Document Register) "
             "and the NDIS certificate of registration. Both are on the Checklist tab as 'Missing'."),
            (rv_row, "Rise version",
             "The same documents, with each policy split into a Policy and a separate Procedure (26 "
             "procedures). The checklist lists the audited document, which holds both.")):
        for col, v in ((1, label), (2, text)):
            # labels take the style of the section they join; the text is plain wrapped prose
            src = ws.cell(rv_row - 1 if (row == rv_row and col == 1) else 26, col)
            dst = ws.cell(row, col)
            dst.value = v
            dst.font, dst.fill = copy(src.font), copy(src.fill)
            dst.alignment = Alignment(wrap_text=True, vertical="top")
    ws.title = "Start here"


# On Harmony's own Staff File Checklist but not yet on the staff file tab. Appended below its last
# row: Team at a glance finds items by name, so the rows above keep their places.
STAFF_FILE_ADDITIONS = [  # (item, who needs it, what to keep, renew, applies)
    ("Letter of offer", "All workers", "Signed letter of offer", "Once", '="Yes"'),
    ("Salary sacrifice arrangement", "Workers who salary sacrifice", "Signed arrangement", "Once", "No"),
    ("Private motor vehicle use agreement", "Drivers", "Signed agreement to use their own car for work, with "
     "the vehicle details", "Once, and when the car changes", '=IF($B$9="Yes","Yes","No")'),
    ("User account", "All workers", "Account set up; removed when they leave", "Removed when they leave", '="Yes"'),
    ("Mobile phone issued", "Workers given a phone", "What was issued and when; returned when they leave",
     "Returned when they leave", "No"),
    ("Credit card issued", "Workers given a card", "Card issued and its limit; returned when they leave",
     "Returned when they leave", "No"),
    ("Office key issued", "Workers given a key", "Key issued; returned when they leave",
     "Returned when they leave", "No"),
]


def add_staff_items(wb, docs):
    from copy import copy
    ws = wb["Staff file (copy per worker)"]
    form_url = next(d["url"] for d in docs if d["title"] == "Staff File Checklist")
    last = ws.max_row
    if ws.cell(last, 2).value != "Unlinked in the NDIS Commission Portal; file kept 7 years":
        raise SystemExit("the staff file tab has changed; check where its rows end before appending")
    cf = next(c for c in ws.conditional_formatting if str(c.sqref) == f"I22:I{last}")
    for n, (item, who, keep, renew, applies) in enumerate(STAFF_FILE_ADDITIONS, last + 1):
        for col in range(1, 12):
            src, dst = ws.cell(last, col), ws.cell(n, col)
            dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
            dst.alignment, dst.number_format = copy(src.alignment), src.number_format
        values = ["Also on Harmony's Staff File Checklist", item, who, keep, "Staff File Checklist", renew, None,
                  None, f'=IF(K{n}="No","Not needed",IF(AND(G{n}="",H{n}=""),"Missing",IF(H{n}="","Held",'
                        f'IF(H{n}<TODAY(),"Expired",IF(H{n}<TODAY()+60,"Renew soon","In date")))))',
                  "Harmony Staff File Checklist." + ("" if applies.startswith("=") else
                                                     " Type Yes in Applies if this worker has one."), applies]
        for col, v in enumerate(values, 1):
            ws.cell(n, col).value = v
        ws.cell(n, 5).hyperlink = form_url
        if not applies.startswith("="):
            ws.cell(n, 11).fill = copy(ws.cell(last, 7).fill)  # the input colour: someone types here
    end = last + len(STAFF_FILE_ADDITIONS)
    for rule in cf.rules:
        new = copy(rule)
        new.formula = [f.replace(f"I22", f"I{last + 1}") for f in rule.formula]
        ws.conditional_formatting.add(f"I{last + 1}:I{end}", new)


def forest_not_plum(wb):
    """The 16 Oct sheet styled its headings plum; the locked Harmony design has no plum."""
    from copy import copy
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.font and c.font.color is not None and c.font.color.rgb == "005C3A52":
                    f = copy(c.font); f.color = FOREST; c.font = f
                if c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb == "005C3A52":
                    c.fill = fill(FOREST)


def main(base, out):
    wb, docs = read_base(base)
    docs += [dict(d) for d in MISSING_DOCS]
    enrich(docs)
    order = groups(docs)
    at = wb.sheetnames.index("All documents")
    del wb["All documents"]
    _, end = write_checklist(wb, order, at)
    write_send_tabs(wb, docs, at + 1)
    write_tracker(wb, at + 1 + len(SEND_TABS) + 1)
    fix_start_here(wb, end)
    add_staff_items(wb, docs)
    forest_not_plum(wb)
    wb.active = 0
    wb.save(out)
    counts = {}
    for d in docs:
        counts[d["file"] or "Not filled in"] = counts.get(d["file"] or "Not filled in", 0) + 1
    print(f"{len(docs)} documents in {len(order)} groups -> {out}")
    print(counts)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
