"""The Everything tab on its own, as a file to hand over: one sheet, every link still opening its file.

On its own there are no other tabs for its formulas to read, so each formula is replaced by the
value it gave on the day of the export (LibreOffice works them out), and the 'Type dates in'
column, which only jumps to those tabs, goes.

    python3 everything_alone.py CHECKLIST.xlsx OUT.xlsx
"""
import datetime as dt
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import everything as ev
import make_workbook as mw

JUMPS = ev.C["Type dates in"]
LAST = get_column_letter(JUMPS - 1)
NO_WORKERS = {"Add workers on Team at a glance": "No workers entered yet"}  # that tab is not in this file


def calculated(path):
    """The workbook's values as Excel would show them; openpyxl writes formulas without results."""
    if not shutil.which("soffice"):
        sys.exit("LibreOffice (soffice) is needed to work out the formulas.")
    with tempfile.TemporaryDirectory() as out:
        run = subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out,
                              str(Path(path).resolve())], capture_output=True, text=True, timeout=600)
        done = Path(out) / Path(path).with_suffix(".xlsx").name
        if not done.exists():  # soffice exits 0 even when it converts nothing
            sys.exit(f"LibreOffice did not work out the formulas:\n{run.stdout}{run.stderr}")
        return openpyxl.load_workbook(done, data_only=True)[ev.TAB]


def main(src, dst):
    if not Path(src).is_file():
        sys.exit(f"No workbook at {src}.")
    today = dt.date.today()
    values = calculated(src)
    wb = openpyxl.load_workbook(src)
    if ev.TAB not in wb.sheetnames:
        sys.exit(f"{src} has no {ev.TAB} tab.")
    ws = wb[ev.TAB]

    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                v = values[c.coordinate].value
                if isinstance(v, str) and v.startswith(("#", "Err:")):
                    sys.exit(f"{c.coordinate} works out to {v}; not writing a file with an error in it.")
                c.value = NO_WORKERS.get(v, v)

    for name in wb.sheetnames:
        if name != ev.TAB:
            del wb[name]
    for name in list(wb.defined_names):
        del wb.defined_names[name]

    bands = [m for m in ws.merged_cells.ranges if m.max_col == JUMPS]
    for m in bands:
        ws.unmerge_cells(m.coord)
    ws.delete_cols(JUMPS)
    for m in bands:
        ws.merge_cells(start_row=m.min_row, start_column=m.min_col, end_row=m.max_row, end_column=JUMPS - 1)
    ws.column_dimensions.pop(get_column_letter(JUMPS), None)
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink is not None and not c.hyperlink.target:
                c.hyperlink = None  # a jump to a tab this file does not have
    ws.auto_filter.ref = f"A{mw.HEADER_ROW}:{LAST}{ws.max_row}"
    if not ws._images:
        mw.add_logo(ws, JUMPS - 1)

    note = ("Under each Practice Standard: every document, every item each worker's file must hold, and every "
            f"dated job, with when it is needed, how often and whether it is in date on "
            f"{today.day} {today:%B %Y}. Dates are kept up to date on the Everything tab of the Harmony Care "
            "checklist; this copy does not change. The 1 and 2 buttons at the top left show the indicators.")
    ws["A3"] = note
    ws["A3"].alignment = mw.WRAP
    ws.merge_cells(f"A3:{LAST}3")  # wrapped inside the page, not running past its edge
    ws.row_dimensions[3].height = mw.height(note, sum(w for _, w in ev.COLS[:JUMPS - 1]))

    left = [c.coordinate for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and (c.value.startswith("=") or "Team at a glance" in c.value)]
    if left:
        sys.exit(f"Still pointing at other tabs: {left[:5]}")
    mw.check_never_named(wb)
    wb.save(dst)
    links = sum(1 for row in ws.iter_rows() for c in row if c.hyperlink is not None)
    print(f"{ws.max_row - mw.HEADER_ROW} rows, {links} links -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
