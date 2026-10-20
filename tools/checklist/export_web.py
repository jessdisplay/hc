"""Export the checklist workbook to site/checklist/ for the checklist page and the team hub.

checklist.json holds every document by Practice Standard; hub.json holds what the hub
adds: the send-out packs, the staff file items and the year planner.

The workbook (made by make_workbook.py, then kept up to date by Harmony) is the
master. This repo is public, so only what the page shows leaves the workbook:
the audit notes and Harmony's rule sources stay behind.

    python3 export_web.py CHECKLIST.xlsx
"""
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

import make_workbook as mw
from send_lists import SEND_TABS

OUT = Path(__file__).resolve().parents[2] / "site" / "checklist" / "checklist.json"
HUB = OUT.with_name("hub.json")
FIELDS = ["no", "title", "url", "type", "file", "when", "how_often", "completed_by", "signed_by",
          "version", "approved_by", "approval_date", "next_review", "review_months", "last_reviewed"]


def iso(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else v


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    year, month = d.year + y, m + 1
    last = (dt.date(year + (month == 12), month % 12 + 1, 1) - dt.timedelta(days=1)).day
    return dt.date(year, month, min(d.day, last))


def read(path):
    ws = openpyxl.load_workbook(path)["Checklist"]
    head = [c.value for c in ws[mw.HEADER_ROW]]
    want = {h: k for h, _, k in mw.COLS}
    docs = []
    for row in ws.iter_rows(min_row=mw.HEADER_ROW + 1):
        if not isinstance(row[0].value, int):
            continue
        cells = dict(zip(head, row))
        d = {want[h]: iso(c.value) for h, c in cells.items() if want.get(h) in FIELDS + ["section"]}
        d["url"] = cells["Document (opens the file)"].hyperlink.target if cells["Document (opens the file)"].hyperlink else None
        d["file"] = None if d["file"] == "Not filled in" else next(k for k, v in mw.FILE_LABEL.items() if v == d["file"])
        status = cells["Status"].value
        d["fixed_status"] = None if (status is None or str(status).startswith("=")) else status
        # Next due follows the workbook's own formula, so the page can judge it against today.
        last, months = cells["Last reviewed (type the date)"].value, d["review_months"]
        if isinstance(last, dt.datetime) and months:
            d["next_due"] = add_months(last.date(), int(months)).isoformat()
        else:
            d["next_due"] = d["next_review"]
        docs.append(d)
    return docs


def structure(docs):
    sections = {s["section_no"]: s for s in mw.STANDARDS["sections"]}
    modules = []
    for module, key, heading, outcome, rows, clause, mine, also in mw.groups(
            [dict(d, section=d.pop("section")) for d in docs]):
        if not modules or modules[-1]["label"] != module:
            modules.append({"label": module, "standards": []})
        modules[-1]["standards"].append({
            "key": key, "heading": heading, "outcome": outcome, "cite": clause,
            "indicators": mw.indicator_lines(rows) if key in sections else [],
            "docs": [d["no"] for d in mine],
            "also": [d["no"] for d in also],
        })
    return modules


def hub(path, docs):
    """The send-out packs as document numbers, plus the staff file and year planner tabs."""
    numbers = {}
    for d in docs:
        if d["fixed_status"] not in mw.NOT_SENT:
            numbers.setdefault(d["title"], []).append(d["no"])
    packs = []
    for t in SEND_TABS:
        groups = []
        for name, entries in t["groups"]:
            items = [{"no": n, "action": action} for title, action in entries for n in numbers[title]]
            groups.append({"name": name, "items": items})
        packs.append({"key": t["tab"], "title": t["title"], "intro": t["intro"], "groups": groups})
    wb = openpyxl.load_workbook(path)
    ws = wb["Staff file (copy per worker)"]
    head = [c.value for c in ws[21]]
    if head[:6] != ["Group", "Item", "Who needs it", "What to keep in the file", "Harmony form", "Renew"]:
        raise SystemExit("the staff file tab's columns have moved")
    staff = [dict(zip(("group", "item", "who", "keep", "form", "renew"), (c.value for c in row[:6])))
             for row in ws.iter_rows(min_row=22) if row[1].value]
    ws = wb["Year planner"]
    if [c.value for c in ws[3]][:2] != ["Due", "What"]:
        raise SystemExit("the year planner tab's columns have moved")
    planner = [{"due": iso(row[0].value), "what": row[1].value}
               for row in ws.iter_rows(min_row=4) if row[1].value]
    return {"packs": packs, "staff": staff, "planner": planner}


def main(path):
    docs = read(path)
    modules = structure([dict(d) for d in docs])
    out = {
        "generated": dt.date.today().isoformat(),
        "standards_source": mw.STANDARDS["source"]["guidelines"],
        "modules": modules,
        "docs": [{k: d.get(k) for k in FIELDS + ["fixed_status", "next_due"]} for d in docs],
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")))
    print(f"{len(docs)} documents, {sum(len(m['standards']) for m in modules)} groups -> {OUT}")
    h = hub(path, docs)
    HUB.write_text(json.dumps(h, ensure_ascii=False, separators=(",", ":")))
    print(f"{sum(len(g['items']) for p in h['packs'] for g in p['groups'])} pack items, "
          f"{len(h['staff'])} staff file items, {len(h['planner'])} planner rows -> {HUB}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
