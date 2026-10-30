"""Everything in one list: each Practice Standard, then every document, every item a worker's
file must hold, and every dated job, with when each is needed and whether it is in date.

Brittany asked for "one for all areas and all forms, when they're needed". The other tabs stay
where dates are typed (documents on Checklist, workers on their own staff file tab, jobs on the
Year planner); this list only reads them, so nothing is typed twice. make_workbook.py writes it
as the 'Everything' tab and make_pdf.py prints it as Part 1, both from collect().
"""
import datetime as dt
import re

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

import make_workbook as mw
from export_web import add_months

TAB = "Everything"
STAFF_TAB, TEAM_TAB, PLANNER_TAB = "Staff file (copy per worker)", "Team at a glance", "Year planner"
STAFF_HEAD = ["Group", "Item", "Who needs it", "What to keep in the file", "Harmony form", "Renew"]
STAFF_FIRST = 22  # the staff file tab's first item row
TEAM_HEAD, TEAM_FIRST, TEAM_LAST = 4, 5, 19  # Team at a glance: header row, worker rows
PLANNER_HEAD = ["Due", "What", "Open", "Why (source)", "Done on (type the date)"]
REPEATS = {"Every month": 1, "Every fortnight": None}  # months between records; None = 14 days

# A staff file item sits under the standard whose indicator asks for it. The rest sit under
# Human Resource Management, whose indicators name pre-employment checks, induction, training,
# supervision and infection control training (17(2) to 17(9)). Each row names the indicator and
# words it must contain, so the build stops if the standard's wording ever moves.
STAFF_PLACE = [  # (item pattern, section, indicator, words the indicator must contain)
    (r"^Medication training", "26", "2", r"workers responsible for administering medication"),
    (r"^Mealtime support training", "26A", "5", r"trained in preparing and providing safe meals"),
    (r"^Emergency and disaster preparedness training", "18A", "9", r"trained in the implementation of the plans"),
    (r"^High intensity competency", "M1-ALL", None, None),
    (r"^Behaviour support plan training", "50", "3", r"skills required to consistently implement"),
    (r"^Restrictive practices training", "50", "5", r"training in the safe use of restrictive practices"),
    (r"^(SIL Worker Training Plan|Induction to each SIL house)", "72D", "1", r"worker development and training"),
    (r"^(SIL Worker Competency|Positive behaviour support|Cultural safety)", "72D", "3",
     r"trained and assessed as competent"),
    (r"^Key personnel suitability", "11", "6", r"suitably qualified"),
]
STAFF_DEFAULT = "17"
WHEN_OF_GROUP = {"Also on Harmony's Staff File Checklist": "When it applies (Harmony's Staff File Checklist)"}

COLS = [  # (header, width)
    ("No.", 6), ("Module", 13), ("Practice Standard", 24), ("What (opens the file)", 46), ("Kind", 15),
    ("For", 20), ("When it is needed", 40), ("Details", 44), ("How often", 24), ("Next due", 13),
    ("Status", 22), ("Type dates in", 20)]
C = {h: i + 1 for i, (h, _) in enumerate(COLS)}
RED = ("Overdue", "Missing")
AMBER = ("Due within 30 days", "No review date on record", "No record yet", "Renew within 60 days",
         "No dates entered", "Add workers on Team at a glance")
GREEN = ("In date", "Done", "Not due yet")
GREY = ("Superseded", "Working copy", "Draft", "Per worker")


def file_id(url):
    m = re.search(r"/d/([\w-]+)|[?&]id=([\w-]+)", url or "")
    return (m.group(1) or m.group(2)) if m else None


def _check_place():
    sections = {s["section_no"]: s for s in mw.STANDARDS["sections"]}
    for pattern, key, item, words in STAFF_PLACE:
        if key in mw.EXTRA:
            continue
        top = [r for r in sections[key]["rows"] if r["parent"] is None and r["item"] == item]
        if not top or not re.search(words, top[0]["text"], re.I):
            raise SystemExit(f"indicator {key}({item}) no longer says {words!r}; recheck STAFF_PLACE {pattern!r}")


def staff_items(wb):
    ws = wb[STAFF_TAB]
    if [c.value for c in ws[STAFF_FIRST - 1]][:6] != STAFF_HEAD:
        raise SystemExit("the staff file tab's columns have moved")
    team = wb[TEAM_TAB]
    tracked = {team.cell(TEAM_HEAD, c).value: get_column_letter(c) for c in range(3, team.max_column + 1)}
    items = []
    for row in ws.iter_rows(min_row=STAFF_FIRST):
        if not row[1].value:
            continue
        group, item, who, keep, form, renew = (c.value for c in row[:6])
        key = next((k for p, k, *_ in STAFF_PLACE if re.search(p, item)), STAFF_DEFAULT)
        items.append({"group": group, "item": item, "who": who, "keep": keep, "renew": renew,
                      "form": None if form in (None, "None", "None in the set") else form,
                      "url": row[4].hyperlink.target if row[4].hyperlink else None,
                      "section": key, "team_col": tracked.get(item)})
    missing = set(tracked) - {i["item"] for i in items}
    if missing:
        raise SystemExit(f"Team at a glance tracks items the staff file no longer has: {sorted(missing)}")
    return items


def planner_jobs(wb, docs):
    ws = wb[PLANNER_TAB]
    if [c.value for c in ws[3]][:5] != PLANNER_HEAD:
        raise SystemExit("the year planner tab's columns have moved")
    by_id = {file_id(d["url"]): d for d in docs if d.get("url")}
    jobs = []
    for row in ws.iter_rows(min_row=4):
        due, what, opens, why, done = (c.value for c in row[:5])
        if not what:
            continue
        if not isinstance(due, dt.datetime) and due not in REPEATS:
            raise SystemExit(f"year planner row {row[0].row}: due {due!r} is neither a date nor {list(REPEATS)}")
        url = row[2].hyperlink.target if row[2].hyperlink else None
        doc = by_id.get(file_id(url))
        if doc is None:
            raise SystemExit(f"year planner row {row[0].row} opens {opens!r}, which is not on the checklist")
        jobs.append({"row": row[0].row, "due": due, "what": what, "opens": opens, "url": url, "why": why,
                     "done": done, "section": doc["section"]})
    return jobs


def by_section(staff, jobs, keys):
    """{section: {"staff": [...], "jobs": [...]}} for every checklist group key."""
    _check_place()
    filed = {k: {"staff": [], "jobs": []} for k in keys}
    for kind, things in (("staff", staff), ("jobs", jobs)):
        for t in things:
            if t["section"] not in filed:
                raise SystemExit(f"{t.get('item') or t['what']!r} is filed under {t['section']}, "
                                 f"which has no group on the checklist")
            filed[t["section"]][kind].append(t)
    return filed


def collect(order, staff, jobs):
    """The checklist groups, each with the staff file items and planner jobs filed under it."""
    filed = by_section(staff, jobs, [g[1] for g in order])
    return [dict(zip(("module", "key", "heading", "outcome", "rows", "clause", "docs", "also"), g), **filed[g[1]])
            for g in order]


def job_status(job, today):
    """The planner job's status on a given day, as the Everything tab's formula works it out."""
    done, due = job["done"], job["due"]
    if due in REPEATS:
        if not isinstance(done, dt.datetime):
            return "No record yet", None
        months = REPEATS[due]
        nxt = add_months(done.date(), months) if months else done.date() + dt.timedelta(days=14)
        return ("Overdue" if nxt < today else "In date"), nxt
    if done:
        return "Done", due.date()
    if due.date() < today:
        return "Overdue", due.date()
    return ("Due within 30 days" if due.date() < today + dt.timedelta(days=30) else "Not due yet"), due.date()



def write(wb, at, groups):
    ws = wb.create_sheet(TAB, at)
    ws.sheet_properties.tabColor = mw.FOREST
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    L = {h: get_column_letter(i) for h, i in C.items()}
    d0, d1 = L["What (opens the file)"], L["How often"]
    span = sum(w for h, w in COLS[C["What (opens the file)"] - 1:C["How often"]])
    head = mw.HEADER_ROW

    mw.add_logo(ws, len(COLS) - 1)
    ws["A2"] = "Everything in one list, by Practice Standard"
    ws["A2"].font = mw.F(bold=True, size=14, color=mw.FOREST)
    ws["A3"] = ("Under each Practice Standard: every document, every item each worker's file must hold, and every "
                "dated job, with when it is needed, the next due date and whether it is in date. Dates are typed "
                "in one place only (documents on 'Checklist', workers on their own staff file tab, jobs on "
                "'Year planner'); this list reads them. The 1 and 2 buttons at the top left show the indicators.")
    ws["A3"].font = mw.F(color=mw.MUTED)
    ws["A4"] = "Highlighted rows go in someone's file:"
    ws["A4"].font = mw.F(color=mw.MUTED)
    for col, who in zip(("E", "F", "G"), ("Participant", "Worker", "Organisation")):
        ws[f"{col}4"] = mw.FILE_LABEL[who] if who != "Organisation" else "The organisation"
        ws[f"{col}4"].fill = mw.fill(mw.TINT[who])
        ws[f"{col}4"].font = mw.F(bold=True)
        ws[f"{col}4"].alignment = Alignment(horizontal="center")
    ws["H4"] = "No colour: policies and guides, read but not filled in"
    ws["H4"].font = mw.F(color=mw.MUTED)

    for i, (h, _) in enumerate(COLS, 1):
        c = ws.cell(head, i, h)
        c.font = mw.F(bold=True, color="FFFFFF")
        c.fill = mw.fill(mw.FOREST)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[head].height = 30

    def put(r, values, tint=None, link=None, bold=False):
        for h, v in values.items():
            c = ws.cell(r, C[h], v)
            c.font, c.alignment, c.border = mw.F(), mw.WRAP, mw.BORDER
            if h == "Next due":
                c.number_format = "d mmm yyyy"
            if h == "Status":
                c.font = mw.F(bold=True)
            if h in ("Module", "Practice Standard"):  # there to filter by; the band above already says it
                c.font = mw.F(size=9, color=mw.MUTED)
        if link:
            t = ws.cell(r, C["What (opens the file)"])
            t.hyperlink = link
            t.font = mw.F(color=mw.FOREST, underline="single", bold=bold)
        if tint:
            for i in range(1, len(COLS) + 1):
                ws.cell(r, i).fill = mw.fill(tint)

    def jump(r, text, target):
        c = ws.cell(r, C["Type dates in"], text)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=target)  # a jump inside the workbook
        c.font, c.alignment = mw.F(color=mw.FOREST, underline="single", size=9), mw.WRAP

    ck_due, ck_status, ck_last = (get_column_letter(mw.COL[k]) for k in ("Next due", "Status", "last_reviewed"))
    team_names = f"'{TEAM_TAB}'!$A${TEAM_FIRST}:$A${TEAM_LAST}"
    r, module_seen = head + 1, None
    for g in groups:
        if g["module"] != module_seen:
            module_seen = g["module"]
            ws.merge_cells(f"A{r}:{L['Type dates in']}{r}")
            c = ws.cell(r, 1, g["module"].upper())
            c.font, c.fill = mw.F(bold=True, size=12, color="FFFFFF"), mw.fill(mw.FOREST)
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[r].height = 22
            r += 1
        ws.cell(r, C["Practice Standard"], g["heading"]).font = mw.F(bold=True, color=mw.FOREST, size=11)
        ws.cell(r, C["Practice Standard"]).alignment = mw.WRAP
        ws.merge_cells(f"{d0}{r}:{d1}{r}")
        text = f"Outcome: {g['outcome']}" + (f"\n({g['clause']})" if g["clause"] else "")
        c = ws.cell(r, C["What (opens the file)"], text)
        c.font, c.alignment = mw.F(italic=True, color=mw.INK), mw.WRAP
        for i in range(1, len(COLS) + 1):
            ws.cell(r, i).border = Border(top=Side(style="medium", color=mw.FOREST))
        ws.row_dimensions[r].height = max(mw.height(g["heading"], 24, 11), mw.height(text, span))
        r += 1
        lines = mw.indicator_lines(g["rows"])
        first = r
        for n, line in enumerate(lines):
            if n == 0:
                ws.cell(r, C["Practice Standard"], "Indicators to be demonstrated").font = mw.F(color=mw.MUTED, size=9)
            ws.merge_cells(f"{d0}{r}:{d1}{r}")
            c = ws.cell(r, C["What (opens the file)"], line)
            c.font, c.alignment = mw.F(color=mw.MUTED, size=9), mw.WRAP
            ws.row_dimensions[r].height = mw.height(line, span, 9)
            r += 1
        if lines:  # folded away: this list is for doing; the Checklist tab shows them open
            ws.row_dimensions.group(first, r - 1, outline_level=1, hidden=True)
        for d in g["also"]:
            ws.cell(r, C["Practice Standard"], "Also meets this standard").font = mw.F(color=mw.MUTED, size=9)
            ws.merge_cells(f"{d0}{r}:{d1}{r}")
            c = ws.cell(r, C["What (opens the file)"], f"No. {d['no']}, {d['title']} (filed under {d['standard']})")
            c.font, c.alignment = mw.F(color=mw.FOREST, underline="single", italic=True), mw.WRAP
            if d["url"]:
                c.hyperlink = d["url"]
            r += 1
        short = g["module"].split(":")[0]
        for d in g["docs"]:
            row = d["row"]  # its row on the Checklist tab, where its dates are typed
            details = (f"Filled in by: {d['completed_by']}. Signed by: {d['signed_by']}".rstrip(".") + "."
                       if d["file"] and d["completed_by"] else None)
            put(r, {"No.": d["no"], "Module": short, "Practice Standard": g["heading"],
                    "What (opens the file)": d["title"], "Kind": d["type"],
                    "For": mw.FILE_LABEL[d["file"]] if d["file"] else "Not filled in",
                    "When it is needed": d["when"], "Details": details, "How often": d["how_often"],
                    "Next due": f'=IF(Checklist!{ck_due}{row}="","",Checklist!{ck_due}{row})',
                    "Status": f"=Checklist!{ck_status}{row}"},
                tint=mw.TINT[d["file"]] if d["file"] else None, link=d["url"], bold=bool(d["file"]))
            jump(r, f"Checklist, No. {d['no']}", f"Checklist!{ck_last}{row}")
            r += 1
        for s in g["staff"]:
            col = s["team_col"]
            if col:
                rng = f"'{TEAM_TAB}'!{col}{TEAM_FIRST}:{col}{TEAM_LAST}"
                due = f'=IF(COUNT({rng})=0,"",MIN({rng}))'
                status = (f'=IF(COUNTA({team_names})=0,"Add workers on Team at a glance",'
                          f'IF(COUNT({rng})=0,"No dates entered",'
                          f'IF(MIN({rng})<TODAY(),"Expired for "&COUNTIF({rng},"<"&TODAY())&" of "&COUNTA({team_names}),'
                          f'IF(MIN({rng})<TODAY()+60,"Renew within 60 days","In date"))))')
            else:
                due, status = None, "Per worker"
            details = s["keep"] + (f" Harmony form: {s['form']}." if s["form"] else "")
            put(r, {"Module": short, "Practice Standard": g["heading"], "What (opens the file)": s["item"],
                    "Kind": "Worker's file", "For": s["who"],
                    "When it is needed": WHEN_OF_GROUP.get(s["group"], s["group"]), "Details": details,
                    "How often": s["renew"], "Next due": due, "Status": status},
                tint=mw.TINT["Worker"], link=s["url"])
            if col:
                jump(r, "Team at a glance", f"'{TEAM_TAB}'!{col}{TEAM_HEAD}")
            else:
                c = ws.cell(r, C["Type dates in"], "Each worker's own tab")
                c.font, c.alignment = mw.F(size=9, color=mw.MUTED), mw.WRAP
            r += 1
        for j in g["jobs"]:
            p, done = j["row"], f"'{PLANNER_TAB}'!E{j['row']}"
            if j["due"] in REPEATS:
                months = REPEATS[j["due"]]
                nxt = f"EDATE({done},{months})" if months else f"({done}+14)"
                due = f'=IF({done}="","",{nxt})'
                status = f'=IF({done}="","No record yet",IF({nxt}<TODAY(),"Overdue","In date"))'
                kind, when = "Regular job", j["due"]
            else:
                when_cell = f"'{PLANNER_TAB}'!A{p}"
                due = f"={when_cell}"
                status = (f'=IF({done}<>"","Done",IF({when_cell}<TODAY(),"Overdue",'
                          f'IF({when_cell}<TODAY()+30,"Due within 30 days","Not due yet")))')
                kind, when = "Dated job", None
            put(r, {"Module": short, "Practice Standard": g["heading"], "What (opens the file)": j["what"],
                    "Kind": kind, "For": "The organisation", "When it is needed": when, "Details": j["why"],
                    "Next due": due, "Status": status},
                tint=mw.TINT["Organisation"], link=j["url"])
            jump(r, "Year planner", f"'{PLANNER_TAB}'!E{p}")
            r += 1

    end = r - 1
    st, nd = L["Status"], L["Next due"]
    top = head + 1
    red, amber, green, grey = mw.fill("F4C7C3"), mw.fill("FCE8B2"), mw.fill("D7EBD5"), mw.fill("E8E4E0")
    rng = f"{st}{top}:{st}{end}"
    rule = lambda values: "OR(" + ",".join(f'{st}{top}="{v}"' for v in values) + ")"  # noqa: E731
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({rule(RED)},LEFT({st}{top},7)="Expired")'], fill=red))
    for values, f_ in ((AMBER, amber), (GREEN, green), (GREY, grey)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[rule(values)], fill=f_))
    rng = f"{nd}{top}:{nd}{end}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({nd}{top}),{nd}{top}<TODAY())"], fill=red))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({nd}{top}),{nd}{top}<TODAY()+30)"],
                                                   fill=amber))
    ws.auto_filter.ref = f"A{head}:{L['Type dates in']}{end}"
    ws.freeze_panes = f"{L['Kind']}{top}"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{head}:{head}"
    ws.print_area = f"A1:{L['Type dates in']}{end}"
    return ws, end
